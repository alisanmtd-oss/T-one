from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from .company_workspace import save_project_verified_values


PROJECT_ID = re.compile(r"^[a-z0-9][a-z0-9._-]{0,79}$")
SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
TABLE_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls", ".json"}
MEDIA_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".zip", ".pdf"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 40 * 1024 * 1024
MAX_REQUEST_BODY_BYTES = 60 * 1024 * 1024

HEADER_ALIASES = {
    "product": {
        "sku",
        "seller_sku",
        "item_sku",
        "product_id",
        "product_name",
        "name",
        "title",
        "item_name",
        "color",
        "size",
        "产品",
        "商品",
        "商品名称",
        "产品名称",
        "颜色",
        "尺码",
    },
    "cost": {
        "cost",
        "unit_cost",
        "purchase_price",
        "landed_cost",
        "product_cost",
        "成本",
        "单位成本",
        "采购价",
    },
    "price": {
        "price",
        "sale_price",
        "selling_price",
        "list_price",
        "price_usd",
        "售价",
        "销售价",
        "零售价",
    },
    "image": {
        "image",
        "images",
        "image_url",
        "image_path",
        "main_image",
        "main_image_url",
        "图片",
        "主图",
        "图片链接",
    },
    "inventory": {
        "inventory",
        "stock",
        "quantity",
        "qty",
        "available_quantity",
        "库存",
        "库存数量",
        "可售库存",
    },
}

CANONICAL_FIELD_ALIASES = {
    "sku": {"sku", "seller_sku", "item_sku", "merchant_sku", "货号", "商家sku"},
    "product_id": {"product_id", "item_id", "商品id", "产品id"},
    "product_name": {
        "product_name",
        "name",
        "title",
        "item_name",
        "商品名称",
        "产品名称",
    },
    "color": {"color", "colour", "颜色"},
    "size": {"size", "尺码", "规格"},
    "cost": {"cost", "unit_cost", "purchase_price", "landed_cost", "product_cost", "成本", "单位成本", "采购价"},
    "price": {"price", "sale_price", "selling_price", "list_price", "price_usd", "售价", "销售价", "零售价"},
    "inventory": {"inventory", "stock", "quantity", "qty", "available_quantity", "库存", "库存数量", "可售库存"},
    "image": {"image", "images", "image_url", "image_path", "main_image", "main_image_url", "图片", "主图", "图片链接"},
}
MAX_CATALOG_ROWS = 100_000


def _load_workspace_project(root: Path, project_id: str) -> dict[str, Any]:
    try:
        workspace = json.loads(
            (root / "config" / "workspace_projects.json").read_text(encoding="utf-8-sig")
        )
    except (OSError, ValueError) as exc:
        raise ValueError("工作区项目配置不可读。") from exc
    project = next(
        (
            item
            for item in workspace.get("projects", [])
            if isinstance(item, dict) and str(item.get("id") or "") == project_id
        ),
        None,
    )
    if project is None:
        raise ValueError("项目不存在。")
    return project


def _decode_file_content(value: Any) -> bytes:
    text = str(value or "")
    if "," in text and text.lstrip().startswith("data:"):
        text = text.split(",", 1)[1]
    try:
        return base64.b64decode(text, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("导入文件不是有效的 Base64 内容。") from exc


def _safe_filename(name: Any) -> str:
    original = Path(str(name or "file")).name
    stem = SAFE_NAME.sub("-", Path(original).stem).strip("-._") or "file"
    suffix = Path(original).suffix.lower()
    return f"{stem[:80]}{suffix}"


def _normal_header(value: Any) -> str:
    return re.sub(r"[\s\-/]+", "_", str(value or "").strip().lower())


def _headers_from_csv(raw: bytes, suffix: str) -> tuple[list[str], int]:
    text = ""
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        return [], 0
    delimiter = "\t" if suffix == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return [], 0
    return [str(value or "").strip() for value in rows[0]], max(0, len(rows) - 1)


def _headers_from_xlsx(raw: bytes) -> tuple[list[str], int]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("需要安装 openpyxl 才能读取 XLSX；也可以先另存为 CSV。") from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(value or "").strip() for value in header_row]
    row_count = max(0, int(sheet.max_row or 0) - 1)
    workbook.close()
    return headers, row_count


def _headers_from_json(raw: bytes) -> tuple[list[str], int]:
    payload = json.loads(raw.decode("utf-8-sig"))
    if isinstance(payload, list):
        first = next((item for item in payload if isinstance(item, dict)), {})
        return [str(key) for key in first], len(payload)
    if isinstance(payload, dict):
        rows = payload.get("rows") if isinstance(payload.get("rows"), list) else []
        if rows:
            first = next((item for item in rows if isinstance(item, dict)), {})
            return [str(key) for key in first], len(rows)
        return [str(key) for key in payload], 1
    return [], 0


def _inspect_table(raw: bytes, suffix: str) -> dict[str, Any]:
    try:
        if suffix in {".csv", ".tsv"}:
            headers, row_count = _headers_from_csv(raw, suffix)
        elif suffix == ".xlsx":
            headers, row_count = _headers_from_xlsx(raw)
        elif suffix == ".json":
            headers, row_count = _headers_from_json(raw)
        else:
            return {
                "headers": [],
                "row_count": 0,
                "detected": {},
                "warning": "旧版 XLS 文件已保存，但自动识别需要先另存为 XLSX 或 CSV。",
            }
    except (OSError, ValueError, json.JSONDecodeError, csv.Error) as exc:
        return {
            "headers": [],
            "row_count": 0,
            "detected": {},
            "warning": f"文件已保存，但自动识别失败：{exc}",
        }

    normalized = {_normal_header(header): header for header in headers if str(header).strip()}
    detected: dict[str, list[str]] = {}
    for field, aliases in HEADER_ALIASES.items():
        matches = [original for normal, original in normalized.items() if normal in aliases]
        if matches:
            detected[field] = matches
    return {
        "headers": headers[:100],
        "row_count": row_count,
        "detected": detected,
        "warning": "",
    }


def _csv_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("商品表编码无法识别，请另存为 UTF-8 CSV。")


def _rows_from_csv(raw: bytes, suffix: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(
        io.StringIO(_csv_text(raw)), delimiter="\t" if suffix == ".tsv" else ","
    )
    rows: list[dict[str, Any]] = []
    for row in reader:
        if len(rows) >= MAX_CATALOG_ROWS:
            raise ValueError("商品表超过 100000 行，请按项目或店铺拆分后导入。")
        if isinstance(row, dict):
            rows.append({str(key or ""): value for key, value in row.items() if key is not None})
    return rows


def _rows_from_xlsx(raw: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("需要安装 openpyxl 才能读取 XLSX；也可以先另存为 CSV。") from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator, ())]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if len(rows) >= MAX_CATALOG_ROWS:
            workbook.close()
            raise ValueError("商品表超过 100000 行，请按项目或店铺拆分后导入。")
        rows.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
        )
    workbook.close()
    return rows


def _rows_from_json(raw: bytes) -> list[dict[str, Any]]:
    payload = json.loads(raw.decode("utf-8-sig"))
    if isinstance(payload, dict) and isinstance(payload.get("rows"), list):
        payload = payload["rows"]
    if isinstance(payload, dict):
        payload = [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON 商品表必须是对象数组，或包含 rows 数组。")
    rows = [item for item in payload if isinstance(item, dict)]
    if len(rows) > MAX_CATALOG_ROWS:
        raise ValueError("商品表超过 100000 行，请按项目或店铺拆分后导入。")
    return [{str(key): value for key, value in item.items()} for item in rows]


def _read_catalog_rows(raw: bytes, suffix: str) -> list[dict[str, Any]]:
    if suffix in {".csv", ".tsv"}:
        return _rows_from_csv(raw, suffix)
    if suffix == ".xlsx":
        return _rows_from_xlsx(raw)
    if suffix == ".json":
        return _rows_from_json(raw)
    return []


def _clean_catalog_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _canonical_header(header: Any) -> str:
    normalized = _normal_header(header)
    for field, aliases in CANONICAL_FIELD_ALIASES.items():
        if normalized in {_normal_header(alias) for alias in aliases}:
            return field
    return ""


def _media_refs_for_row(
    canonical: dict[str, Any], media_records: list[dict[str, Any]], row_count: int
) -> list[str]:
    by_name: dict[str, str] = {}
    for item in media_records:
        name = str(item.get("name") or "").lower()
        path = str(item.get("path") or "")
        if name and path:
            by_name[name] = path
            by_name[Path(name).stem] = path
    matches: list[str] = []
    raw_image = str(canonical.get("image") or "")
    for token in re.split(r"[,;|\n]+", raw_image):
        key = Path(token.strip()).name.lower()
        match = by_name.get(key) or by_name.get(Path(key).stem)
        if match and match not in matches:
            matches.append(match)
    sku = str(canonical.get("sku") or "").strip().lower()
    if sku:
        for item in media_records:
            if sku in Path(str(item.get("name") or "")).stem.lower():
                path = str(item.get("path") or "")
                if path and path not in matches:
                    matches.append(path)
    if not matches and row_count == 1 and len(media_records) == 1:
        matches.append(str(media_records[0].get("path") or ""))
    return [item for item in matches if item]


def _catalog_dedupe_key(canonical: dict[str, Any], row_number: int) -> str:
    sku = str(canonical.get("sku") or "").strip()
    if sku:
        return f"sku:{sku.lower()}"
    product_id = str(canonical.get("product_id") or "").strip()
    if product_id:
        return f"product_id:{product_id.lower()}"
    identity = "|".join(
        str(canonical.get(field) or "").strip().lower()
        for field in ("product_name", "color", "size")
    )
    return f"variant:{identity}" if identity.strip("|") else f"row:{row_number}"


def _build_catalog(
    project_id: str,
    table_reference: str,
    rows: list[dict[str, Any]],
    media_records: list[dict[str, Any]],
) -> dict[str, Any]:
    dynamic_fields: set[str] = set()
    records_by_key: dict[str, dict[str, Any]] = {}
    duplicate_count = 0
    row_count = len(rows)
    for row_number, row in enumerate(rows, start=2):
        canonical: dict[str, Any] = {}
        attributes: dict[str, Any] = {}
        for header, raw_value in row.items():
            value = _clean_catalog_value(raw_value)
            if value in (None, ""):
                continue
            field = _canonical_header(header)
            if field:
                canonical.setdefault(field, value)
                continue
            attribute = _normal_header(header)
            if attribute:
                attributes[attribute] = value
                dynamic_fields.add(attribute)
        if not canonical and not attributes:
            continue
        key = _catalog_dedupe_key(canonical, row_number)
        media_refs = _media_refs_for_row(canonical, media_records, row_count)
        current = records_by_key.get(key)
        if current is not None:
            duplicate_count += 1
            current["source_rows"].append(row_number)
            for field, value in canonical.items():
                current["fields"].setdefault(field, value)
            for field, value in attributes.items():
                current["attributes"].setdefault(field, value)
            for path in media_refs:
                if path not in current["media_refs"]:
                    current["media_refs"].append(path)
            continue
        record_id = hashlib.sha256(f"{project_id}:{key}".encode("utf-8")).hexdigest()[:20]
        records_by_key[key] = {
            "record_id": record_id,
            "dedupe_key": key,
            "fields": canonical,
            "attributes": attributes,
            "media_refs": media_refs,
            "source_rows": [row_number],
        }
    records = list(records_by_key.values())
    product_keys = {
        str(record["fields"].get("product_name") or record["fields"].get("product_id") or record["fields"].get("sku") or record["record_id"])
        .strip()
        .lower()
        for record in records
    }
    return {
        "schema_version": 1,
        "project_id": project_id,
        "source_table_ref": table_reference,
        "normalization_mode": "deterministic_schema",
        "ai_enrichment_status": "awaiting_platform_target",
        "canonical_fields": sorted(CANONICAL_FIELD_ALIASES),
        "dynamic_fields": sorted(dynamic_fields),
        "source_extension_fields": sorted(dynamic_fields),
        "field_origin": "imported_source_columns_not_platform_schema",
        "platform_schema_fields": [],
        "product_count": len(product_keys),
        "variant_count": len(records),
        "duplicate_count": duplicate_count,
        "records": records,
    }


def save_product_data_intake(root: Path, payload: dict[str, Any]) -> dict[str, Any]:
    project_id = str(payload.get("project_id") or "").strip().lower()
    if not PROJECT_ID.fullmatch(project_id):
        raise ValueError("请选择一个有效项目组。")
    project = _load_workspace_project(root, project_id)
    files = payload.get("files") if isinstance(payload.get("files"), list) else []
    manual = payload.get("manual") if isinstance(payload.get("manual"), dict) else {}
    if not files and not any(str(value or "").strip() for value in manual.values()):
        raise ValueError("请导入商品表、图片，或填写至少一项商品数据。")
    if len(files) > 20:
        raise ValueError("一次最多导入 20 个文件。")

    prepared_files: list[dict[str, Any]] = []
    total_bytes = 0
    table_count = 0
    for item in files:
        if not isinstance(item, dict):
            raise ValueError("导入文件记录格式无效。")
        kind = str(item.get("kind") or "").strip().lower()
        if kind not in {"table", "media"}:
            raise ValueError("导入文件缺少有效类型。")
        filename = _safe_filename(item.get("name"))
        suffix = Path(filename).suffix.lower()
        allowed = TABLE_EXTENSIONS if kind == "table" else MEDIA_EXTENSIONS
        if suffix not in allowed:
            raise ValueError(f"不支持的文件类型：{filename}")
        raw = _decode_file_content(item.get("content_base64"))
        if len(raw) > MAX_FILE_BYTES:
            raise ValueError(f"{filename} 超过 20MB，请拆分后再导入。")
        total_bytes += len(raw)
        if total_bytes > MAX_TOTAL_BYTES:
            raise ValueError("本次导入总大小超过 40MB，请分批导入。")
        if kind == "table":
            table_count += 1
        prepared_files.append(
            {
                "kind": kind,
                "name": filename,
                "suffix": suffix,
                "raw": raw,
                "size_bytes": len(raw),
                "sha256": hashlib.sha256(raw).hexdigest(),
            }
        )
    if table_count > 1:
        raise ValueError("一次只能导入一张商品表；请先合并表格或分项目导入。")

    stamp = datetime.now(timezone.utc).astimezone().strftime("%Y%m%d-%H%M%S-%f")
    intake_id = f"{stamp}-{uuid4().hex[:8]}"
    target_dir = root / "data" / "product_intake" / project_id / intake_id
    target_dir.mkdir(parents=True, exist_ok=True)
    saved: list[dict[str, Any]] = []
    table_inspection: dict[str, Any] = {}
    table_reference = ""
    media_references: list[str] = []
    table_raw = b""
    table_suffix = ""

    for item in prepared_files:
        kind = str(item["kind"])
        filename = str(item["name"])
        suffix = str(item["suffix"])
        raw = bytes(item["raw"])
        target = target_dir / f"{uuid4().hex[:8]}-{filename}"
        target.write_bytes(raw)
        relative = target.relative_to(root).as_posix()
        record = {
            "kind": kind,
            "name": filename,
            "path": relative,
            "size_bytes": int(item["size_bytes"]),
            "sha256": str(item["sha256"]),
        }
        if kind == "table" and not table_reference:
            table_reference = relative
            table_raw = raw
            table_suffix = suffix
            table_inspection = _inspect_table(raw, suffix)
            record["inspection"] = table_inspection
        elif kind == "media":
            media_references.append(relative)
        saved.append(record)

    detected = (
        table_inspection.get("detected")
        if isinstance(table_inspection.get("detected"), dict)
        else {}
    )
    manual_product = str(manual.get("product") or "").strip()
    manual_cost = str(manual.get("cost") or "").strip()
    manual_price = str(manual.get("price") or "").strip()
    manual_inventory = str(manual.get("inventory") or "").strip()

    values: dict[str, Any] = {}
    if manual_product or detected.get("product"):
        values["product_catalog_source_ref"] = manual_product or (
            f"{table_reference} · {', '.join(detected.get('product') or [])}"
        )
    if manual_cost or detected.get("cost"):
        values["product_cost_source_ref"] = manual_cost or (
            f"{table_reference} · {', '.join(detected.get('cost') or [])}"
        )
    if manual_price or detected.get("price"):
        values["pricing_source_ref"] = manual_price or (
            f"{table_reference} · {', '.join(detected.get('price') or [])}"
        )
    if media_references or detected.get("image"):
        values["product_image_source_ref"] = (
            ", ".join(media_references[:8])
            if media_references
            else f"{table_reference} · {', '.join(detected.get('image') or [])}"
        )
    if manual_inventory or detected.get("inventory"):
        inventory_reference = manual_inventory or (
            f"{table_reference} · {', '.join(detected.get('inventory') or [])}"
        )
        values["inventory_source_ref"] = inventory_reference
        values["capacity_source_ref"] = inventory_reference

    project_snapshot = save_project_verified_values(root, project_id, values)
    media_records = [item for item in saved if item.get("kind") == "media"]
    catalog: dict[str, Any] = {
        "schema_version": 1,
        "project_id": project_id,
        "normalization_mode": "manual_reference_only",
        "ai_enrichment_status": "awaiting_platform_target",
        "canonical_fields": sorted(CANONICAL_FIELD_ALIASES),
        "dynamic_fields": [],
        "source_extension_fields": [],
        "field_origin": "manual_reference_only_not_platform_schema",
        "platform_schema_fields": [],
        "product_count": 0,
        "variant_count": 0,
        "duplicate_count": 0,
        "records": [],
    }
    if table_raw and table_suffix != ".xls":
        try:
            catalog = _build_catalog(
                project_id,
                table_reference,
                _read_catalog_rows(table_raw, table_suffix),
                media_records,
            )
        except (OSError, ValueError, json.JSONDecodeError, csv.Error) as exc:
            warning = str(table_inspection.get("warning") or "")
            table_inspection["warning"] = (
                f"{warning}；商品记录规范化失败：{exc}" if warning else f"商品记录规范化失败：{exc}"
            )
    catalog_path = target_dir / "catalog.json"
    catalog_path.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    manifest = {
        "schema_version": 1,
        "created_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "project_id": project_id,
        "project_name": str(project.get("name") or project_id),
        "files": saved,
        "manual_fields": sorted(key for key, value in manual.items() if str(value or "").strip()),
        "detected": detected,
        "row_count": int(table_inspection.get("row_count") or 0),
        "warning": str(table_inspection.get("warning") or ""),
        "catalog_path": catalog_path.relative_to(root).as_posix(),
        "product_count": int(catalog.get("product_count") or 0),
        "variant_count": int(catalog.get("variant_count") or 0),
        "duplicate_count": int(catalog.get("duplicate_count") or 0),
        "dynamic_fields": catalog.get("dynamic_fields") or [],
        "source_extension_fields": catalog.get("source_extension_fields") or [],
        "catalog_field_origin": str(catalog.get("field_origin") or ""),
        "platform_schema_field_count": 0,
        "external_connector_call_count": 0,
        "external_action_count": 0,
        "ai_enrichment_status": catalog.get("ai_enrichment_status"),
        "missing_fields": project_snapshot.get("owner_blockers") or [],
    }
    (target_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    manifest_ref = (target_dir / "manifest.json").relative_to(root).as_posix()
    catalog_ref = catalog_path.relative_to(root).as_posix()
    context_refs = {
        "attachment_manifest_ref": manifest_ref,
        "source_refs": [manifest_ref, catalog_ref],
        "business_object_refs": [
            {
                "object_type": "product_intake_batch",
                "object_id": intake_id,
                "version_ref": catalog_ref,
                "source_ref": manifest_ref,
            }
        ],
    }
    return {
        "ok": True,
        "project_id": project_id,
        "saved_file_count": len(saved),
        "row_count": manifest["row_count"],
        "detected": detected,
        "missing_fields": manifest["missing_fields"],
        "owner_required_count": int(project_snapshot.get("owner_required_count") or 0),
        "owner_provided_count": int(project_snapshot.get("owner_provided_count") or 0),
        "owner_missing_count": int(project_snapshot.get("owner_missing_count") or 0),
        "warning": manifest["warning"],
        "intake_id": intake_id,
        "manifest_path": manifest_ref,
        "catalog_path": catalog_ref,
        "context_refs": context_refs,
        "product_count": manifest["product_count"],
        "variant_count": manifest["variant_count"],
        "duplicate_count": manifest["duplicate_count"],
        "dynamic_fields": manifest["dynamic_fields"],
        "source_extension_fields": manifest["source_extension_fields"],
        "catalog_field_origin": manifest["catalog_field_origin"],
        "platform_schema_field_count": 0,
        "external_connector_call_count": 0,
        "external_action_count": 0,
        "ai_enrichment_status": manifest["ai_enrichment_status"],
    }
