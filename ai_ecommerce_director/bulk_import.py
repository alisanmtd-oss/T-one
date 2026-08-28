from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .parser import ALIAS_TO_FIELD, FIELD_ALIASES, RECORD_TYPES, canonical_record_type, normalize_fields
from .scoring import score_record
from .storage import append_record, read_records, write_records
from .types import IntakeRecord


SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
DEFAULT_CONFIDENCE_THRESHOLD = 0.7

REQUIRED_FIELDS = {
    "money_case": {"product_name"},
    "influencer_cooperation": {"influencer_handle"},
    "hot_link": {"url"},
    "supply_chain": {"product_type"},
    "culture_label": {"english_label"},
    "holiday_opportunity": {"holiday_name"},
    "tk_rule_update": {"summary"},
    "store_metric_snapshot": {"marketplace"},
}

EXPORT_FIELDS = {
    "money_case": [
        "product_name",
        "platform",
        "selling_price",
        "landed_cost",
        "sample_cost",
        "influencer_cost",
        "ad_spend",
        "total_revenue",
        "total_profit",
        "roi",
        "lifecycle",
        "influencer_type",
        "content_angle",
        "hook",
        "main_visual_action",
        "winning_reason",
        "failure_reason",
        "lessons",
    ],
    "influencer_cooperation": [
        "influencer_handle",
        "platform",
        "profile_url",
        "follower_count",
        "avg_views",
        "engagement_rate",
        "quote_price",
        "sample_cost",
        "product_name",
        "video_posted",
        "orders",
        "total_revenue",
        "total_profit",
        "roi",
    ],
    "hot_link": [
        "source_platform",
        "url",
        "screenshot_path",
        "product_name",
        "selling_price",
        "visible_growth_signal",
    ],
    "supply_chain": [
        "product_type",
        "gender",
        "color",
        "size_range",
        "weight_gsm",
        "inventory_qty",
        "unit_cost",
        "sample_cost",
        "production_days",
        "shipping_days",
    ],
    "culture_label": [
        "english_label",
        "chinese_meaning",
        "audience_identity",
        "buying_motivation",
        "content_angle",
    ],
    "holiday_opportunity": [
        "holiday_name",
        "date_or_period",
        "content_angle",
        "product_name",
    ],
    "tk_rule_update": [
        "rule_type",
        "url",
        "summary",
        "risk_level",
        "immediate_action",
    ],
    "store_metric_snapshot": [
        "marketplace",
        "store_id",
        "store_name",
        "currency",
        "metric_date",
        "snapshot_start",
        "snapshot_end",
        "report_type",
        "settlement_id",
        "total_revenue",
        "net_revenue",
        "settlement_amount",
        "order_count",
        "unit_count",
        "ad_spend",
        "refund_amount",
        "refund_count",
        "return_count",
        "platform_fee",
        "commission_fee",
        "transaction_fee",
        "shipping_fee",
        "tax_amount",
        "adjustment_amount",
        "reserve_amount",
    ],
}


@dataclass(slots=True)
class ColumnMapping:
    source_column: str
    canonical_field: str | None
    confidence: float

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_column": self.source_column,
            "canonical_field": self.canonical_field or "unmapped",
            "confidence": self.confidence,
        }


@dataclass(slots=True)
class ImportSummary:
    file_path: str
    record_type: str
    rows_read: int = 0
    created: int = 0
    skipped: int = 0
    duplicates: int = 0
    flagged: int = 0
    missing_required: dict[str, int] = field(default_factory=dict)
    mapping: list[ColumnMapping] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "file_path": self.file_path,
            "record_type": self.record_type,
            "rows_read": self.rows_read,
            "created": self.created,
            "skipped": self.skipped,
            "duplicates": self.duplicates,
            "flagged": self.flagged,
            "missing_required": self.missing_required,
            "mapping": [item.to_dict() for item in self.mapping],
            "errors": self.errors,
        }


def import_table(
    root: Path,
    file_path: Path,
    record_type: str,
    worksheet: str | None = None,
    explicit_mapping: dict[str, str] | None = None,
    confidence_threshold: float = DEFAULT_CONFIDENCE_THRESHOLD,
    update_on_duplicate: bool = False,
) -> ImportSummary:
    canonical_type = canonical_record_type(record_type)
    if canonical_type not in RECORD_TYPES:
        raise ValueError(f"Unsupported record type: {record_type}. Supported: {', '.join(sorted(RECORD_TYPES))}")

    summary = ImportSummary(file_path=str(file_path), record_type=canonical_type)
    rows = read_table(file_path, worksheet)
    summary.rows_read = len(rows)
    headers = list(rows[0].keys()) if rows else []
    mapping = map_columns(headers, explicit_mapping or {}, confidence_threshold)
    summary.mapping = mapping

    existing = read_records(root, canonical_type)
    existing_by_key = {dedupe_key(record.record_type, record.fields): record for record in existing}
    updated_existing = list(existing)
    updated_by_id = {record.record_id: index for index, record in enumerate(updated_existing)}

    for row_index, row in enumerate(rows, start=2):
        fields, confidence = row_to_fields(row, mapping)
        missing = sorted(REQUIRED_FIELDS.get(canonical_type, set()) - {key for key, value in fields.items() if value not in {"", None}})
        if missing:
            summary.skipped += 1
            for field_name in missing:
                summary.missing_required[field_name] = summary.missing_required.get(field_name, 0) + 1
            continue

        record_key = dedupe_key(canonical_type, fields)
        if record_key in existing_by_key:
            summary.duplicates += 1
            if update_on_duplicate and existing_by_key[record_key].record_id in updated_by_id:
                existing_record = existing_by_key[record_key]
                existing_record.fields.update(fields)
                existing_record.confidence = max(existing_record.confidence, confidence)
                existing_record.metadata["last_bulk_import_row"] = row_index
                score, notes = score_record(canonical_type, existing_record.fields)
                existing_record.score = score
                existing_record.director_notes = notes
                updated_existing[updated_by_id[existing_record.record_id]] = existing_record
            continue

        review_status = "pending" if confidence < confidence_threshold else "approved"
        score, notes = ({}, ["低置信度记录已进入人工复核队列，暂不进入评分。"]) if review_status == "pending" else score_record(canonical_type, fields)
        record = IntakeRecord(
            record_type=canonical_type,
            fields=fields,
            raw_text=json.dumps(row, ensure_ascii=False),
            source="bulk_import",
            score=score,
            director_notes=notes,
            review_status=review_status,
            confidence=confidence,
            metadata={"file_path": str(file_path), "row_number": row_index},
        )
        append_record(root, record)
        existing_by_key[record_key] = record
        summary.created += 1
        if review_status == "pending":
            summary.flagged += 1

    if update_on_duplicate and updated_existing:
        write_records(root, canonical_type, updated_existing)

    return summary


def export_table(root: Path, record_type: str, file_path: Path) -> int:
    canonical_type = canonical_record_type(record_type)
    if canonical_type not in RECORD_TYPES:
        raise ValueError(f"Unsupported record type: {record_type}. Supported: {', '.join(sorted(RECORD_TYPES))}")

    records = read_records(root, canonical_type)
    base_fields = EXPORT_FIELDS.get(canonical_type, [])
    observed_fields = sorted({field for record in records for field in record.fields})
    headers = list(dict.fromkeys([*base_fields, *observed_fields, "record_id", "source", "review_status", "confidence"]))

    file_path.parent.mkdir(parents=True, exist_ok=True)
    with file_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for record in records:
            row = {field: record.fields.get(field, "") for field in headers}
            row["record_id"] = record.record_id
            row["source"] = record.source
            row["review_status"] = record.review_status
            row["confidence"] = record.confidence
            writer.writerow(row)
    return len(records)


def read_table(file_path: Path, worksheet: str | None = None) -> list[dict[str, Any]]:
    if not file_path.exists():
        raise FileNotFoundError(f"Import file does not exist: {file_path}")

    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file extension {extension}. Supported extensions: {supported}")

    if extension == ".csv":
        return read_csv(file_path)
    if extension == ".xlsx":
        return read_xlsx(file_path, worksheet)
    return read_xls(file_path, worksheet)


def read_csv(file_path: Path) -> list[dict[str, Any]]:
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            return []
        return [{key: clean_value(value) for key, value in row.items() if key is not None} for row in reader]


def read_xlsx(file_path: Path, worksheet: str | None = None) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading .xlsx files requires openpyxl.") from exc

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[worksheet] if worksheet else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    table: list[dict[str, Any]] = []
    for raw_row in rows[1:]:
        row = {
            headers[index]: clean_value(value)
            for index, value in enumerate(raw_row)
            if index < len(headers) and headers[index]
        }
        if any(value not in {"", None} for value in row.values()):
            table.append(row)
    return table


def read_xls(file_path: Path, worksheet: str | None = None) -> list[dict[str, Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("Reading .xls files requires xlrd. Convert the file to .xlsx or .csv, or install xlrd.") from exc

    workbook = xlrd.open_workbook(file_path)
    sheet = workbook.sheet_by_name(worksheet) if worksheet else workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return []
    headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
    table: list[dict[str, Any]] = []
    for row_index in range(1, sheet.nrows):
        row = {
            headers[col]: clean_value(sheet.cell_value(row_index, col))
            for col in range(sheet.ncols)
            if headers[col]
        }
        if any(value not in {"", None} for value in row.values()):
            table.append(row)
    return table


def map_columns(
    headers: list[str],
    explicit_mapping: dict[str, str],
    confidence_threshold: float,
) -> list[ColumnMapping]:
    mappings: list[ColumnMapping] = []
    for header in headers:
        if header in explicit_mapping:
            canonical = canonical_field(explicit_mapping[header])
            mappings.append(ColumnMapping(header, canonical, 1.0 if canonical else 0.0))
            continue

        canonical, confidence = infer_field(header)
        if confidence < confidence_threshold:
            canonical = None
        mappings.append(ColumnMapping(header, canonical, confidence))
    return mappings


def infer_field(header: str) -> tuple[str | None, float]:
    normalized = normalize_header(header)
    if normalized in ALIAS_TO_FIELD:
        return ALIAS_TO_FIELD[normalized], 1.0
    for field_name, aliases in FIELD_ALIASES.items():
        if normalized == normalize_header(field_name):
            return field_name, 1.0
        for alias in aliases:
            alias_norm = normalize_header(alias)
            if normalized == alias_norm:
                return field_name, 1.0
            if alias_norm and alias_norm in normalized:
                return field_name, 0.85
    return None, 0.0


def row_to_fields(row: dict[str, Any], mapping: list[ColumnMapping]) -> tuple[dict[str, Any], float]:
    raw_fields: dict[str, Any] = {}
    confidences: list[float] = []
    for item in mapping:
        if not item.canonical_field:
            continue
        value = row.get(item.source_column)
        if value in {"", None}:
            continue
        raw_fields[item.canonical_field] = value
        confidences.append(item.confidence)
    fields = normalize_fields(raw_fields)
    confidence = min(confidences) if confidences else 0.0
    return fields, confidence


def dedupe_key(record_type: str, fields: dict[str, Any]) -> tuple[Any, ...]:
    if record_type == "hot_link":
        return (record_type, fields.get("url"))
    if record_type == "influencer_cooperation":
        return (record_type, fields.get("profile_url") or fields.get("influencer_handle"), fields.get("product_name"))
    if record_type == "supply_chain":
        return (
            record_type,
            fields.get("product_type"),
            fields.get("color"),
            fields.get("size_range"),
            fields.get("weight_gsm"),
        )
    if record_type == "culture_label":
        return (record_type, fields.get("english_label"))
    if record_type == "holiday_opportunity":
        return (record_type, fields.get("holiday_name"), fields.get("date_or_period"))
    if record_type == "tk_rule_update":
        return (record_type, fields.get("summary"), fields.get("rule_type"))
    if record_type == "store_metric_snapshot":
        return (
            record_type,
            fields.get("marketplace") or fields.get("source_platform"),
            fields.get("store_id") or fields.get("store_name"),
            fields.get("settlement_id") or fields.get("metric_date") or fields.get("snapshot_end"),
            fields.get("report_type"),
        )
    return (
        record_type,
        fields.get("url"),
        fields.get("product_name"),
        fields.get("platform"),
        fields.get("roi"),
    )


def canonical_field(name: str) -> str | None:
    normalized = normalize_header(name)
    if normalized in ALIAS_TO_FIELD:
        return ALIAS_TO_FIELD[normalized]
    if name in FIELD_ALIASES:
        return name
    return None


def normalize_header(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_").replace("-", "_")


def clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value
